"""A location explorer driven entirely by the cards' capability blocks.

One `Location` = one ordered, two-faced deck (from the workbook) + a `Board` (from
mechanisms/card_effects.py). Turning, dispositions, role passives, gates and upkeep
are all read from the blocks — no per-location logic. The same class drives Reef
Gardens, Crash Flats and Kelp Cathedral. The deck cursor persists across turns, so
a location resumes exactly where the previous turn left it.

Gates come in two shapes, both read from a card's `back.gate`:
  * mid-deck marker (Reef Caverns, Engine Core): turning STOPS here; the deep opens
    only if the icon requirement is met from the present team's hands + this location.
  * entry gate on the location card itself (Kelp Cathedral, card 1): the team cannot
    even enter unless the requirement is met.
"""
from __future__ import annotations
import os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, ".."))
for _rt in (os.path.join(HERE, "..", "..", "spec", "runtime"),
            os.path.join(HERE, "..", "spec", "runtime")):
    if os.path.isdir(os.path.join(_rt, "gamespec")):
        sys.path.insert(0, _rt); break

from gamespec.engine import iter_design_objects
from mechanisms import card_effects as CE

GAME = os.path.join(HERE, "..", "game")
WB = os.path.join(HERE, "..", "components", "components.xlsx")
EFFECT_FILES = {
    "Reef Gardens":   os.path.join(GAME, "reef_cards.gamespec.md"),
    "Crash Flats":    os.path.join(GAME, "crash_flats_cards.gamespec.md"),
    "Kelp Cathedral": os.path.join(GAME, "kelp_cathedral_cards.gamespec.md"),
}
_NUM = re.compile(r"_(\d+)_")


def load_all_effects():
    eff = {}
    for f in EFFECT_FILES.values():
        eff.update(CE.load_effects(iter_design_objects(f)))
    return eff


def load_recipes(location):
    """Aphelion Salvage recipes available at a location (the s* cards): craftable during
    the team phase. Returns their card ids in deck order."""
    import openpyxl
    wb = openpyxl.load_workbook(WB, data_only=True)
    ws = wb["Cards"]; rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[0]]
    gi = lambda n: hdr.index(n)
    ID, LOC, KW = gi("ID"), gi("Location"), gi("Keywords")
    out = []
    for r in rows[1:]:
        if not r or r[ID] is None or str(r[LOC]) != location:
            continue
        kws = str(r[KW] or "").lower()
        if "aphelion salvage" in kws or "outpost upgrade" in kws:
            out.append(str(r[ID]))
    return out


def load_deck(location):
    """Ordered, numbered turn-deck for one location (recipes s*/cascades c* excluded —
    they enter play by crafting/triggers, not by being turned). Types + keywords come
    from the workbook; all rules come from the capability blocks."""
    import openpyxl
    wb = openpyxl.load_workbook(WB, data_only=True)
    ws = wb["Cards"]; rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[0]]
    gi = lambda n: hdr.index(n)
    ID, NM, CT, LOC, KW = [gi(x) for x in ["ID", "Name", "Card Type", "Location", "Keywords"]]
    deck = []
    for r in rows[1:]:
        if not r or r[ID] is None or str(r[LOC]) != location:
            continue
        m = _NUM.search(str(r[ID]))
        if not m:                                   # s1 / s3 / c2 → not turned
            continue
        deck.append({"id": str(r[ID]), "name": str(r[NM]), "type": str(r[CT]).lower(),
                     "pos": int(m.group(1)),
                     "keywords": [w.strip().lower() for w in re.split(r"[,\n]", str(r[KW] or "")) if w.strip()]})
    deck.sort(key=lambda c: c["pos"])
    return deck


class Location:
    def __init__(self, name, effects):
        self.name = name
        self.effects = effects
        self.deck = load_deck(name)
        self.recipes = load_recipes(name)
        self.cursor = 0
        self.board = CE.Board(effects=effects)
        self.entered = False
        self.status = "unexplored"          # unexplored | open | gate-closed | entry-denied | exhausted

    # ---- helpers --------------------------------------------------------
    def _disposition(self, e):
        for blk in (e.get("back", {}), e.get("front", {})):
            if isinstance(blk, dict):
                if blk.get("disposition"):
                    return blk["disposition"]
                for step in CE._as_list(blk.get("on_reveal")):
                    if isinstance(step, dict) and step.get("disposition"):
                        return step["disposition"]
        return "resolve_now"

    def _play(self, card, trace, cultivator):
        cid = card["id"]; front = self.effects.get(cid, {}).get("front", {})
        ip = CE.InPlay(cid, front=dict(front))
        ip.front["_card_type"] = card["type"]
        ip.front["_keywords"] = card["keywords"]
        self.board.in_play[cid] = ip
        CE.apply_effects(self.effects.get(cid, {}).get("back", {}).get("on_reveal"), self.board, ip)
        trace.append(("play", card["name"], "%s, tracks %s" % (card["type"], ip.tracks)))
        # Systems Ecologist's Cultivator: once per turn, +1 Pop to a flora as it lands
        if card["type"] == "flora" and not cultivator[0]:
            for p in self.board.players:
                act = next((x["action"] for x in p.passives if "action" in x), None)
                if act and act.get("id") == "cultivator":
                    CE.apply_effect({"change_track": act["change_track"]}, self.board, ip)
                    cultivator[0] = True
                    trace.append(("action", "%s: Cultivator" % p.name,
                                  "%s.population -> %d" % (card["name"], ip.tracks.get("population", 0))))
                    break
        return ip

    def _gate(self, card, trace, allow_hypothetical=False):
        """Evaluate a back.gate requirement against the present team's icons."""
        req = self.effects.get(card["id"], {}).get("back", {}).get("gate", {}).get("requires", {})
        avail = CE.available_icons(self.board)
        met = all(avail.get(k, 0) >= v for k, v in req.items())
        trace.append(("GATE", card["name"], "need %s · have %s · %s"
                      % (req, avail, "OPEN" if met else "CLOSED")))
        return met, req, avail

    # ---- the turn -------------------------------------------------------
    def explore(self, players_here, trace, dilemma_choice="protect", cultivator=None, assist=False):
        self.board.players = players_here
        cultivator = cultivator if cultivator is not None else [False]
        while self.cursor < len(self.deck):
            card = self.deck[self.cursor]; cid = card["id"]
            e = self.effects.get(cid, {}); back = e.get("back", {}); front = e.get("front", {})
            disp = self._disposition(e)
            has_gate = "gate" in back

            # entry gate on the location card itself (Kelp): must pay to enter
            if has_gate and disp == "play_at_location" and not self.entered:
                met, req, avail = self._gate(card, trace)
                if not (met or assist):
                    self.status = "entry-denied"; return
                if assist and not met:
                    trace.append(("assist", "(hypothetical)", "entry forced to exercise %s's cards" % self.name))
                self.entered = True; self.status = "open"  # paid; fall through to play it

            # mid-deck expedition marker (Reef Caverns / Engine Core): turning stops
            if has_gate and disp != "play_at_location":
                met, req, avail = self._gate(card, trace)
                if not (met or assist):
                    self.status = "gate-closed"; return
                if assist and not met:
                    trace.append(("assist", "(hypothetical)", "gate forced open to exercise the deep"))
                self.status = "open"
                self.cursor += 1; continue                # descend

            # department reveal → assign the present players their role + passives
            if any(isinstance(s, dict) and "each_player_chooses" in s for s in CE._as_list(back.get("on_reveal"))):
                rc = next(s for s in CE._as_list(back.get("on_reveal")) if "reveal_cards" in s)
                positions = rc["reveal_cards"]["positions"]
                trace.append(("reveal", card["name"], "roles at positions %s" % positions))
                for pi, p in enumerate(players_here):
                    if pi >= len(positions):
                        break
                    role = next((c for c in self.deck if c["pos"] == positions[pi]), None)
                    if not role:
                        continue
                    p.role_id = role["id"]
                    for pas in CE._as_list(self.effects.get(role["id"], {}).get("back", {}).get("passive")):
                        p.passives.append(pas)
                    trace.append(("role", "%s = %s" % (p.name, role["name"]), _passive_desc(self.effects.get(role["id"], {}))))
                # skip past the role cards we just assigned
                self.cursor += 1 + len(positions); continue

            if disp == "play_at_location":
                self.entered = True
                self._play(card, trace, cultivator)
            elif disp == "to_player_hand":
                (players_here[0] if players_here else self.board.players[0]).hand.append(cid)
                trace.append(("to_hand", card["name"], "to a player's hand"))
            elif disp == "set_beside_deck":
                trace.append(("aside", card["name"],
                              "beside the deck" + (" (conditional gain)" if back.get("conditional_gain") else "")))
            elif "dilemma" in front:
                opts = front["dilemma"]["options"]
                opt = next((o for o in opts if o["id"] == dilemma_choice), opts[0])
                trace.append(("dilemma", card["name"], "chose '%s'" % opt["label"]))
                CE.apply_effects(opt["do"], self.board, None)
                for d in opt["do"]:
                    if "provides_icon" in d:
                        pinfo = d["provides_icon"]
                        self.board.tucked_icons[pinfo["icon"]] = self.board.tucked_icons.get(pinfo["icon"], 0) + pinfo.get("amount", 1)
            else:
                trace.append(("resolve", card["name"], "read / resolved"))
            self.cursor += 1
        if self.cursor >= len(self.deck):                 # ran out of cards to turn
            self.status = "exhausted"

    def resolve_gains(self, trace):
        for cid, e in self.effects.items():
            for cg in CE._as_list(e.get("back", {}).get("conditional_gain")):
                if cid in self.board.gained or not self.board.players:
                    continue
                if CE.eval_condition(cg.get("when"), self.board, None):
                    self.board.gained.append(cid)                 # taken once, to a present player
                    self.board.players[0].hand.append(cid)
                    trace.append(("gain", CE._short(cid), "conditional gain met"))

    def upkeep(self, trace):
        for ip in list(self.board.in_play.values()):
            if ip.removed:
                continue
            up = self.effects.get(ip.id, {}).get("front", {}).get("upkeep")
            if up is None:
                continue
            before, unresolved = CE.run_upkeep(up, self.board, ip)
            tag = "upkeep-unresolved" if unresolved else "upkeep"
            if before != ip.tracks or unresolved:
                trace.append((tag, CE._short(ip.id), "%s -> %s" % (before, ip.tracks)))

    def craft_phase(self, trace):
        """Team phase: build whatever salvage the location can now afford. Recipes only
        become active once the exploration has produced the resources — the Wreckage
        Field's Materials pay for the Oxygen Tank, and so on. Loops so a build that frees
        nothing new simply stops."""
        if not self.entered:
            return
        made = True
        while made:
            made = False
            for cid in self.recipes:
                who = self.board.players[0] if self.board.players else None
                res = CE.try_craft(cid, self.board, who)
                if res:
                    trace.append(("craft", CE._short(cid), res))
                    made = True

    def site_icons(self):
        """Icons that stay with the location: in-play cards' provided icons + tucked."""
        icons = {}
        for ip in self.board.in_play.values():
            if ip.removed:
                continue
            for pi in CE._as_list(ip.front.get("provides_icon")):
                if CE.eval_condition(pi.get("when"), self.board, ip):
                    icons[pi["icon"]] = icons.get(pi["icon"], 0) + pi.get("amount", 1)
        for icon, amt in self.board.tucked_icons.items():
            icons[icon] = icons.get(icon, 0) + amt
        return icons

    def icons(self):
        return CE.available_icons(self.board)


def player_icons(player, effects):
    """Portable icons a player carries anywhere: role passives + recipes in hand."""
    icons = {}
    for pas in player.passives:
        if "provides_icon" in pas:
            pi = pas["provides_icon"]
            icons[pi["icon"]] = icons.get(pi["icon"], 0) + pi.get("amount", 1)
    for cid in player.hand:
        for pi in CE._as_list(effects.get(cid, {}).get("front", {}).get("provides_icon")):
            icons[pi["icon"]] = icons.get(pi["icon"], 0) + pi.get("amount", 1)
    return icons


def _passive_desc(e):
    out = []
    for p in CE._as_list(e.get("back", {}).get("passive")):
        if "provides_icon" in p:
            out.append("passive %d %s icon" % (p["provides_icon"].get("amount", 1), p["provides_icon"]["icon"]))
        if "action" in p:
            out.append("passive action: %s" % p["action"].get("id"))
    return "; ".join(out) or "—"
