from pathlib import Path
import json,hashlib,re,shutil,datetime
import openpyxl

ROOT=Path(r'C:\AnalogueGames\aphelion-deep')
ROOTS=[Path(r'C:\AnalogueGames\analogue-projects\GameProjects\aphelion_deep'),Path(r'C:\AnalogueGames\TheAphelionDeep'),Path(r'C:\Art and Design\The Aphelion Deep'),Path(r'C:\AnalogueGames\analogue-projects\GameProjects\_shared_worlds\abyssfall')]
def save(path,obj):
 p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf8')
def text(path,s):
 p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf8')
inventory=[];seen={}
for n,root in enumerate(ROOTS):
 for p in sorted(root.rglob('*')):
  if not p.is_file():continue
  rel=p.relative_to(root);digest=hashlib.sha256(p.read_bytes()).hexdigest()
  excluded=any(x in rel.parts for x in ('.git','__pycache__','.obsidian','.pytest_cache')) or (n==0 and rel.parts[0]=='spec')
  dest=None
  if excluded:disp='excluded';reason='Generated cache or vendored generic GameSpec tooling; original preserved'
  elif digest in seen:disp='duplicate';dest=seen[digest];reason='Identical SHA256'
  else:
   asset=p.suffix.lower() in ('.png','.jpg','.jpeg','.svg','.af','.pdf')
   dest=str(Path('assets' if asset else 'archive')/f'source-{n+1}'/rel).replace('\\','/')
   out=ROOT/dest;out.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,out);seen[digest]=dest
   disp='asset' if asset else 'archived';reason='Preserved original bytes; selection report identifies active derivatives'
  inventory.append(dict(source_id=f'origin-{n+1}:{rel.as_posix()}',original_path=str(p),relative_path=rel.as_posix(),sha256=digest,bytes=p.stat().st_size,modified_utc=datetime.datetime.fromtimestamp(p.stat().st_mtime,datetime.timezone.utc).isoformat(),disposition=disp,reason=reason,destination=dest,drive_file_id=None))
save(Path('provenance/inventory.json'),inventory)
def rows(path,sheet):
 w=openpyxl.load_workbook(path,read_only=True,data_only=True);s=w[sheet];r=list(s.iter_rows(max_row=min(s.max_row,1000),values_only=True));return [dict(zip(r[0],x)) for x in r[1:] if x[0] is not None]
v3=rows(ROOTS[1]/'abyssfall_cards_v3.xlsx','Playtest Cards')
cards=rows(ROOTS[0]/'source/components/components.xlsx','Cards')
lookup={(str(r['ID Number']),r['Location']):r for r in v3}
for c in cards:
 r=lookup[(str(c['Source ID']),c['Location'])]
 for key in ('Name','Keywords','Back Text','Front Text','Track Name','Designer Notes'):c[key]=r[key]
 for key,v in [('Track Art Ref','Track Art'),('Back Icon Ref','Back Icon'),('Front Icon Ref','Front Icon'),('Background Art Ref','Background')]:c[key]=r[v]
 c['Original Status']=c['Status'];c['Status']='working';c['Deck Order']=r['ID Number'];c['V3 Source']='source-2/abyssfall_cards_v3.xlsx#Playtest Cards';c['Quantity Note']=re.search(r'Copies: (\d+)',r['Designer Notes'] or '').group(0) if re.search(r'Copies: (\d+)',r['Designer Notes'] or '') else ''
 c['Source Fingerprint']=hashlib.sha256(json.dumps({k:c[k] for k in ('ID','Name','Keywords','Back Text','Front Text','Track Name')},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
save(Path('src/aphelion_deep/data/cards.json'),cards)
w=openpyxl.load_workbook(ROOTS[1]/'abyssfall_cards_v3.xlsx',read_only=True,data_only=True)
pacing='\n\n'.join(str(r[0]) for r in w['Turn Pacing'].iter_rows(values_only=True) if r[0])
text(Path('design/legacy-pacing.md'),pacing)
outline=(ROOTS[1]/'megagame_design_outline.md').read_text(encoding='utf8')
sections={};head='Introduction'
for line in outline.splitlines():
 if line.startswith('## '):head=line[3:];sections[head]=[]
 sections.setdefault(head,[]).append(line)
def pick(names):return '\n\n'.join('\n'.join(sections[n]) for n in names)
docs={}
def doc(slug,title,body,private=False):
 s=f'# {title}\n\nConsolidated from the supplied Aphelion Deep design sources. Working design; unresolved passages retain their original status. The tutorial amendments in Core Rules govern the declared three turn slice.\n\n'+body
 text(Path('design')/(slug+'.md'),s);docs[slug]={'title':title,'text':s,'private':private}
doc('overview','The Aphelion Deep Overview',pick(['Working Title','Core Premise','Design Goals','Player Experience Arc','Shared Mechanical Core','Summary']))
agreed='''[PT-ID:ad-tutorial-rules] [WORKING] The tutorial is three players over three turns, with Exploration followed by Upkeep. Ordered location decks persist. Full colony Command, automatic reactor escalation and challenge injection are outside this tutorial.

[PT-ID:ad-source-authority] [WORKING] Unambiguous v3 printed card text overrides simulator behavior. The following agreed amendments are authoritative for the tutorial.

[PT-ID:ad-track-limits] [WORKING] Reef rations starts at 5. All card tracks use the artwork endpoints 0 and 5. Other setup values follow printed cards. Reactor starts at 0 and advances only through card effects; its actual value supplies Reactor thresholds.

[PT-ID:ad-cultivator] [WORKING] Cultivator is optional, once per turn during Exploration, targeting one explicitly selected Flora.

[PT-ID:ad-feeding] [WORKING] Spotted Grazers consume Flora. Carrier Crabs consume other Fauna. Full feeding payment is required before any population is consumed; otherwise apply starvation. Calcite Spires Population references mean Abundance.

[PT-ID:ad-upkeep] [WORKING] Resolve every revealed location once per turn, even if unoccupied. Only present players eat or craft. Optional crafting, harvesting and taking ecosystem cards into hand occur at most once per stated phase unless otherwise specified. Mandatory upkeep still resolves.

[PT-ID:ad-crafting] [WORKING] Use printed crafting costs. Fabricator discounts one chosen resource by one, with a zero floor. Salvage Manifest Knowledge applies only to crafting Aphelion Salvage items, never expedition gates.

[PT-ID:ad-cut-free] [WORKING] Cut it free gives Giant Lungweed Stalks 2 Population and the location 1 Materials, then removes the dilemma. Recover it and Document and leave remain unresolved.

[PT-ID:ad-route] [WORKING] Turn 1 Fabrication Lead visits Crash Flats and both Ecologists visit Reef Gardens; choose Restart and Protect. Turn 2 Fabrication Lead and Lead Ecologist visit Kelp Cathedral; Systems Ecologist stays at Reef. Choose Cut it free. Turn 3 all three regroup at Reef and attempt Reef Caverns with their actual accumulated capabilities. Missing choices and unsupported actions block execution; gates are never forced.
'''
doc('core-rules','The Aphelion Deep Core Rules',agreed+'\n## Full colony design outside the tutorial\n\n'+pick(['Turn Structure','Outpost System','Depth and Outpost Scaling','Permaculture and Closed Economy','Reactor Clock','Event System','Negotiation and Command Authority']))
doc('departments','The Aphelion Deep Departments',pick(['Departments','Combined Role Model for a 30 Player Game','Role and Philosophy Interactions']))
for kind in ('engineering','ecology','medical'):
 doc(kind+'-briefs',f'The Aphelion Deep {kind.title()} Briefs',(ROOTS[1]/f'abyssfall_{kind}_briefs.md').read_text(encoding='utf8'))
doc('philosophies','The Aphelion Deep Philosophies',pick(['Philosophies','Tension Between Cooperation and Conflict']))
locbody=pick(['Map and Information','Exploration Tiers and Biomes','Suggested Dominant Question by Tier'])
for loc in ('Reef Gardens','Crash Flats','Kelp Cathedral'):
 locbody+='\n\n## '+loc+' current deck\n\n'
 for c in cards:
  if c['Location']==loc:locbody+=f"### {c['Name']}\n\nStable card ID {c['ID']}. Canonical editable faces are in Card Catalogue.\n\n{c['Back Text'] or ''}\n\n{c['Front Text'] or ''}\n\n"
doc('locations','The Aphelion Deep Locations',locbody)
lore=(ROOTS[1]/'abyssfall_drowned_brief.md').read_text(encoding='utf8');m=re.search(r'^## 2[\.\s—]',lore,re.M)
first=re.search(r'^## 1[\.\s—]',lore,re.M)
assert m and first
doc('player-lore','The Aphelion Deep Player Lore',lore[first.start():m.start()])
doc('facilitator-lore','The Aphelion Deep Facilitator Lore',lore[m.start():]+'\n\n'+pick(['The Drowned','Megafauna']),True)
doc('facilitator-guide','The Aphelion Deep Facilitator Guidance',pick(['Important Design Risks','Next Design Steps'])+'\n\n## Original playtest pacing notes\n\nThese notes describe the earlier playtest configuration and are preserved in full as historical facilitation context. The current tutorial route is in Core Rules.\n\n'+pacing,True)
doc('card-design','The Aphelion Deep Card Design Reference',(ROOTS[1]/'abyssfall_card_design.md').read_text(encoding='utf8'),True)
doc('sensory-vocabulary','The Aphelion Deep Sensory Vocabulary',(ROOTS[1]/'abyssfall_sensory_vocabulary.md').read_text(encoding='utf8'))
doc('deeper-knowledge','The Aphelion Deep Deeper Knowledge Drafts',(ROOTS[1]/'abyssfall_tier4_knowledge_cards.md').read_text(encoding='utf8'),True)
doc('structured-design','The Aphelion Deep Structured Design Annex','July GameSpec extensions are preserved below. These declarations describe the broader game and do not override the tutorial amendments.\n\n'+'\n\n'.join(p.read_text(encoding='utf8') for p in sorted((ROOTS[0]/'source/game').glob('*.md')) if not p.name.endswith('_cards.gamespec.md'))+'\n\n## Abyssfall world dependency\n\n'+'\n\n'.join(p.read_text(encoding='utf8') for p in sorted(ROOTS[3].rglob('*.md'))),True)
save(Path('provenance/staged-documents.json'),docs)
headers=list(cards[0]);save(Path('provenance/card-table.json'),[headers]+[[c.get(h) if c.get(h) is not None else '' for h in headers] for c in cards])
text(Path('design/source-selection.md'),'''# Source selection

The v3 workbook supplies the current 46 card definitions, preserving both faces verbatim. Stable IDs, quantity, visibility, collections and card types come from July components.xlsx. Copy counts in designer notes remain separate because the structured Quantity field is 1. Approved tutorial amendments are explicit in Core Rules, not silent edits to the historical printed text.

The June prose is retained in full through section imports and source archives. July GameSpec modules extend the structural design but are not complete replacements for prose; their claimed imports and simulator results do not establish authority over printed text. Full v3 pacing is recovered rather than the truncated components heading. The earlier 178 row card catalogue and v2 variants remain archived outside the active manifest.

The later editable Playmat.af and full map exports are retained alongside the older labelled AD_Mat.png. There is no established one-to-one mapping from the labelled art to the twenty prose locations. Maps and prose are not silently renamed.

The original runtime and effect definitions remain archived as implementation evidence. The game plugin derives executable mappings from the current canonical capture plus the approved amendments, and rejects changed sources until mappings are reviewed.

provenance/inventory.json records every source, hash, modification date, disposition and preserved destination. Dates support the comparison but do not determine content authority. Generic vendored GameSpec tooling and caches are excluded with reasons; all original directories remain untouched.
''')
text(Path('design/gaps.md'),'''# Unresolved content

P0 Tutorial source mapping: any source change requires review and refreshed bindings before simulation. Missing or unsupported script choices must stop the run.

P1 Department completion: Command, Exploration and Corporate/Data have concepts and role outlines but no separate detailed briefs. Engineering personal goals include an unfinished "1) .". Ecology playtest covers two of five roles; Engineering one of five.

P1 Outcomes: Power Question Strip for parts has an unnamed abundance target. Kelp Recover it and Document and leave remain ???. Lost Signal has no resolved front. No tutorial implementation may silently choose these outcomes.

P1 Full colony: Command allocations, reactor escalation, events, Wraith injection, deep expeditions and endgame need integration and balancing beyond this tutorial.

P1 Map naming: artwork and prose use conflicting twenty-location rosters. Kelp Cathedral/Cathedrals also differs; no identity equivalence has been asserted.

P2 Quantities: structured Quantity 1 conflicts with Copies 3 for tanks/welders and Copies 2 for generators. Preserve both; tutorial does not build beyond one of each declared recipe.

P2 Stale references: A.E.V./E.S.S. Aphelion, Reef Garden/Gardens, Energy/Power, Wreck/Wreckage, Watcher card 21 (actually 20), and missing design.md references need editorial decisions. Tutorial uses the July capability mapping Energy=Power and Wreck=Wreckage, retaining original text.

P2 Ecosystem transport: carrying track values, exact portable ecosystem benefits and movement timing need general rules beyond the tutorial policy that leaves ecosystem cards in place. Control question resolution remains human facilitated.

P2 Balance: three turns are a rules demonstration, not evidence of a balanced thirty player megagame. Older pacing is retained as historical context.
''')
art=['Sunlit Reef','Shoal Channel','Coral Gardens','Seagrass Meadows','Shelf Edge','Wreckage Field','Kelp Forest','Drowned Arches','Kelp Corridors','Broken Trenchline','Silt Temple','Overhang','Vent Fields','Trench Pass','Abyssal Trench','Drowned Spines','Hall of Records','Crimson Foundry','Frozen Rift','Reactor Depths']
prose=re.findall(r'^#### \d+\. (.+)$',outline,re.M)
save(Path('provenance/map-reconciliation.json'),{'artwork_labels':art,'prose_locations':prose,'status':'Unresolved. Lists are independent; positions do not assert equivalence.'})
text(Path('.gitignore'),'.venv/\n__pycache__/\n.pytest_cache/\n.playtrace/*.db\n.staging/\n*.egg-info/\n')
print(json.dumps({'inventoried':len(inventory),'unique_preserved':len(seen),'cards':len(cards),'documents':len(docs),'locations':len(prose)}))
