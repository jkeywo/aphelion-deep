from pathlib import Path
import json,re,subprocess,sys
from docx import Document
from docx.shared import Pt,Inches,RGBColor
root=Path(r'C:\AnalogueGames\aphelion-deep')
out=root/'.staging';out.mkdir(exist_ok=True)
docs=json.loads((root/'provenance/staged-documents.json').read_text(encoding='utf8'))
for slug,item in docs.items():
 d=Document();section=d.sections[0];section.top_margin=section.bottom_margin=Inches(.75)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6)
 for name in ['Title','Heading 1','Heading 2','Heading 3']:
  d.styles[name].font.color.rgb=RGBColor(0,0,0)
 code=False
 for line in item['text'].splitlines():
  if line.startswith('```'):code=not code;continue
  if not line.strip():continue
  if not code and line.startswith('#'):
   m=re.match(r'^(#+)\s+(.*)',line)
   if m:
    level=len(m[1]);title=re.sub(r'[^\w\s]',' ',m[2]);d.add_paragraph(title,'Title' if level==1 else f'Heading {min(3,level-1)}');continue
  p=d.add_paragraph(line)
  if code:
   p.paragraph_format.space_after=Pt(0)
   for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8)
 d.save(out/(slug+'.docx'))
 sanitizer=Path(r'C:\Users\jkeyw\.codex\plugins\cache\openai-primary-runtime\documents\26.1009.11318\skills\documents\scripts\google_docs_title_sanitize.py')
 subprocess.run([sys.executable,str(sanitizer),str(out/(slug+'.docx')),'--out',str(out/(slug+'-import.docx'))],check=True,stdout=subprocess.DEVNULL)
print('Staged',len(docs),'documents')
